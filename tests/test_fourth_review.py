"""Regressions for the fourth review pass: each test names the defect it pins."""

import io
import time

import httpx
import pytest
import respx
from fastapi.testclient import TestClient

from app import netguard
from app.ingest import IngestError, parse_bytes
from app.jobs import output_header
from app.main import _attachment, app
from app.pipeline import fetch as F
from app.pipeline.dns_check import DnsResult
from app.pipeline.discovery.score import score_url, select
from app.pipeline.extract.emails import extract_page, fix_tld, unglue_sentence
from app.pipeline.extract.normalize_html import normalize_html
from app.pipeline.filter_rank import classify
from app.pipeline.htmlscan import bodies, iter_blocks, strip_blocks
from app.pipeline.parked import check_ns
from app.pipeline.run import HOMEPAGE_LADDER, ladder_for
from app.ratelimit import CircuitBreaker, CircuitOpen, FetchGate


# --- linear block scanning ---------------------------------------------------


def test_unclosed_script_tags_are_linear_not_quadratic():
    """160 KB of bare `<script>` took 15 s per pass on the worker's event loop."""
    html = "<script>" * 20_000
    started = time.perf_counter()
    F.strip_non_jsonld_scripts(html)
    extract_page(html, "https://acme.com/")
    normalize_html(html)
    assert time.perf_counter() - started < 2.0


def test_block_scanner_matches_the_regex_on_well_formed_html():
    html = ('<p>a</p><SCRIPT type="text/javascript">x</SCRIPT><script>y</script >'
            '<svg><path/></svg><svgfoo>keep</svgfoo>')
    assert bodies(html, "script") == ["x", "y"]
    assert [b.attrs for b in iter_blocks(html, "script")] == [' type="text/javascript"', ""]
    assert strip_blocks(html, "svg") == '<p>a</p><SCRIPT type="text/javascript">x</SCRIPT>' \
                                        '<script>y</script > <svgfoo>keep</svgfoo>'


def test_unclosed_block_runs_to_the_end_of_the_document():
    assert strip_blocks("<p>a</p><style>b{}", "style") == "<p>a</p> "


# --- scripts, blocked pages ------------------------------------------------------


def test_inline_script_with_an_address_survives_stripping():
    html = '<html><body><script>var e="info@acme.com";</script><script>var x=1</script>hi</body></html>'
    out = F.strip_non_jsonld_scripts(html)
    assert "info@acme.com" in out and "var x=1" not in out
    found, text, _ = extract_page(out, "https://acme.com/")
    assert [c.email for c in found] == ["info@acme.com"]
    assert found[0].method == "script"
    assert "info@acme.com" not in text          # never leaks into the visible text


def test_recaptcha_markup_is_not_a_bot_challenge():
    page = ('<html><head><script src="https://www.google.com/recaptcha/api.js"></script>'
            '<style>.captcha{}</style></head><body class="g-recaptcha">info@acme.com</body></html>')
    assert not F.looks_blocked(200, page)
    assert F.looks_blocked(503, "<html><head><title>Just a moment...</title></head></html>")
    assert F.looks_blocked(200, "<html><body>Please complete the captcha to continue</body></html>")
    assert F.looks_blocked(200, '<html><body><form id="cf-chl-widget"></form></body></html>')


# --- fabrications ----------------------------------------------------------------


@pytest.mark.parametrize("text", [
    "Sign up with Gmail. Find us at acmeroofing dot com",
    "Join our Mailchimp list. Shop online at acmeroofing dot com",
])
def test_cue_word_must_be_a_whole_word(text):
    assert "@" not in normalize_html(text)


def test_real_cue_still_deobfuscates():
    assert normalize_html("Contact: bob at acmeroofing dot com") == "Contact: bob@acmeroofing.com"


def test_sentence_glued_after_a_real_tld_is_trimmed():
    assert unglue_sentence("info@acme.com.Call") == "info@acme.com"
    assert unglue_sentence("bob@acme.co.uk") == "bob@acme.co.uk"
    assert unglue_sentence("x@acme.com.au") == "x@acme.com.au"
    found, _, _ = extract_page("<p>Email info@acme.com.Call us today</p>", "https://acme.com/")
    assert [c.email for c in found] == ["info@acme.com"]


def test_one_letter_tail_is_not_glued_text():
    assert fix_tld("info@acme.usa") is None
    assert fix_tld("info@acme.comCall") == "info@acme.com"


# --- discovery fetches the published URL ------------------------------------


def test_fetch_url_keeps_scheme_www_case_and_query():
    c = score_url("http://www.acme.com/Contact-Us.aspx?lang=en#top", "Contact")
    assert c.url == "https://acme.com/contact-us.aspx"
    assert c.target == "http://www.acme.com/Contact-Us.aspx?lang=en"


def test_dedupe_still_collapses_spellings():
    cands = [score_url("http://www.acme.com/Contact"), score_url("https://acme.com/contact/")]
    assert len(select(cands, limit=4)) == 1


# --- homepage ladder --------------------------------------------------------------


def test_ladder_skips_rungs_dns_ruled_out():
    assert ladder_for(DnsResult(resolves=True, a_domain=True, a_www=False)) == (
        ("https", False), ("http", False))
    assert ladder_for(DnsResult(resolves=True, a_domain=False, a_www=True)) == (("https", True),)
    assert ladder_for(DnsResult(resolves=True, a_domain=True, a_www=True)) == HOMEPAGE_LADDER
    assert ladder_for(None) == HOMEPAGE_LADDER


# --- ranking, parked, breaker -----------------------------------------------------


def test_free_provider_is_never_a_sibling_brand():
    own, _, free, rank = classify("jane@outlook.com", "outlookdental.com")
    assert not own and free and rank == 3
    assert classify("info@outlookdental.com", "outlookdental.com")[3] == 2


def test_enquiry_is_a_role_address():
    assert classify("enquiry@acme.com", "acme.com")[3] == 2
    assert classify("inquiry@acme.com", "acme.com")[3] == 2


def test_parking_nameserver_markers_match_whole_labels():
    assert not check_ns(["ns1.jordan.com"])
    assert not check_ns(["ns1.bodisoft.net"])
    assert check_ns(["ns1.dan.com"])
    assert check_ns(["ns2.bodis.com"])
    assert check_ns(["NS1.SEDOPARKING.COM."])


def test_half_open_admits_only_probes_many_probes_in_flight():
    b = CircuitBreaker("jina", open_seconds=0.0, min_samples=1, probes=5)
    b.record(False)
    b.record(False)
    assert b.state == "half_open"
    admitted = 0
    for _ in range(50):
        try:
            b.check()
            admitted += 1
        except CircuitOpen:
            pass
    assert admitted == 5


# --- private-address guard ---------------------------------------------------------


@pytest.mark.parametrize("ip", ["10.0.0.1", "127.0.0.1", "169.254.169.254", "fd12::1",
                                "::1", "::ffff:10.0.0.1", "64:ff9b::a00:1", "2002:a00:1::1"])
def test_private_and_embedded_addresses_are_not_public(ip):
    import ipaddress

    assert not netguard.is_public(ipaddress.ip_address(ip))


@respx.mock
async def test_direct_fetch_refuses_a_host_resolving_to_a_private_address(monkeypatch):
    monkeypatch.setattr(F.settings, "fetch_guard_private_addresses", True)

    async def resolve(host, port):
        raise netguard.PrivateAddress(host)

    monkeypatch.setattr(netguard, "resolve_public", resolve)
    route = respx.get("https://internal.acme.com/robots.txt").mock(
        return_value=httpx.Response(200, text="Sitemap: x"))
    gate = FetchGate(rate_per_minute=None, global_concurrency=2, per_domain_concurrency=2)
    with pytest.raises(F.FetchError) as e:
        await F.HttpxFetcher(gate).get_text("https://internal.acme.com/robots.txt")
    assert e.value.reason == "private_address" and not e.value.retryable
    assert not route.called


@respx.mock
async def test_every_redirect_hop_is_guarded(monkeypatch):
    monkeypatch.setattr(F.settings, "fetch_guard_private_addresses", True)
    seen = []

    async def resolve(host, port):
        seen.append((host, port))
        if host.startswith("[") or host == "fd12::1":
            raise netguard.PrivateAddress(host)
        return ["203.0.113.1"]

    monkeypatch.setattr(netguard, "resolve_public", resolve)
    respx.get("https://acme.com/").mock(
        return_value=httpx.Response(302, headers={"location": "http://[fd12::1]:8080/admin"}))
    inner = respx.get("http://[fd12::1]:8080/admin").mock(
        return_value=httpx.Response(200, text="secret@internal"))
    gate = FetchGate(rate_per_minute=None, global_concurrency=2, per_domain_concurrency=2)
    with pytest.raises(F.FetchError) as e:
        await F.HttpxFetcher(gate).get_html("https://acme.com/")
    assert e.value.reason == "private_address"
    assert seen == [("acme.com", 443), ("fd12::1", 8080)]
    assert not inner.called


@respx.mock
async def test_redirects_are_still_followed_to_public_hosts():
    respx.get("http://acme.com/").mock(
        return_value=httpx.Response(301, headers={"location": "https://www.acme.com/"}))
    respx.get("https://www.acme.com/").mock(
        return_value=httpx.Response(200, headers={"content-type": "text/html"},
                                    text="<html>info@acme.com</html>"))
    gate = FetchGate(rate_per_minute=None, global_concurrency=2, per_domain_concurrency=2)
    r = await F.HttpxFetcher(gate).get_html("http://acme.com/")
    assert r.final_url == "https://www.acme.com/" and "info@acme.com" in r.html


@respx.mock
async def test_too_many_redirects_is_permanent():
    respx.get(url__startswith="https://acme.com/").mock(
        side_effect=lambda req: httpx.Response(302, headers={"location": str(req.url) + "x"}))
    gate = FetchGate(rate_per_minute=None, global_concurrency=2, per_domain_concurrency=2)
    with pytest.raises(F.FetchError) as e:
        await F.HttpxFetcher(gate).get_html("https://acme.com/")
    assert e.value.reason == "redirect_loop"


def test_direct_fetches_do_not_share_the_jina_gate():
    jina_gate = FetchGate(rate_per_minute=500, global_concurrency=2, per_domain_concurrency=2)
    f = F.JinaFetcher(jina_gate)
    assert f._text.gate is not jina_gate
    assert f._text.gate.bucket is None


# --- uploads -------------------------------------------------------------------------


def test_wide_csv_is_refused_not_expanded():
    raw = (",".join(f"c{i}" for i in range(2000)) + "\n" + "acme.com\n" * 10).encode()
    with pytest.raises(IngestError) as e:
        parse_bytes(raw, "wide.csv")
    assert e.value.code == "too_many_columns"


def test_trailing_blank_columns_do_not_count():
    raw = ("website" + "," * 600 + "\nacme.com\n").encode()
    parsed = parse_bytes(raw, "x.csv")
    assert parsed.headers == ["website"]


def test_xlsx_width_comes_from_the_header_not_the_sheet_dimensions():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "website"
    ws["A2"] = "acme.com"
    ws["XFD5"] = "far away"
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_bytes(buf.getvalue(), "wide.xlsx")
    assert parsed.headers == ["website"]
    assert all(len(r) == 1 for r in parsed.rows)


def test_oversized_csv_cell_is_invalid_file_not_500():
    raw = ("website\n" + "x" * 200_000 + "\n").encode()
    with pytest.raises(IngestError) as e:
        parse_bytes(raw, "x.csv")
    assert e.value.code == "invalid_file"


def test_upload_body_is_not_read_before_the_key_is_checked(monkeypatch):
    calls = []
    monkeypatch.setattr("app.main.parse_bytes", lambda *a: calls.append(a) or (_ for _ in ()).throw(RuntimeError))
    client = TestClient(app)
    r = client.post("/jobs/preview", files={"file": ("x.csv", b"website\nacme.com\n", "text/csv")})
    assert r.status_code == 401
    assert calls == []


def test_declared_oversized_upload_is_refused_by_content_length(api_key):
    client = TestClient(app)
    r = client.post("/jobs/preview", headers={"X-API-Key": api_key,
                                              "Content-Type": "multipart/form-data; boundary=x",
                                              "Content-Length": str(10**9)})
    assert r.status_code == 413


def test_header_names_get_the_formula_guard():
    assert output_header(['=HYPERLINK("http://x")', "website"])[:2] == \
        ["'=HYPERLINK(\"http://x\")", "website"]


@pytest.mark.parametrize("name,expected_ascii", [
    ("日本.csv", "__"),
    ('a"b.xlsx', "a_b"),
    ("leads.csv", "leads"),
    (None, "results"),
])
def test_download_name_is_latin1_safe(name, expected_ascii):
    header = _attachment(name, "_extracted.csv")
    header.encode("latin-1")
    assert f'filename="{expected_ascii}_extracted.csv"' in header
    assert "filename*=UTF-8''" in header
