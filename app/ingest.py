"""Section 14: file parsing and website-column detection.

Parses CSV/TSV/TXT/XLSX, keeps every original column and the original row order, and
works out which column holds the website. Detection is header match first, then content
scoring, then an email-derived fallback.
"""

import csv
import io
import logging
import re
from dataclasses import dataclass, field

from app.pipeline.normalize import FREE_EMAIL_LABELS, NormalizedInput, normalize_input

log = logging.getLogger("email_extractor.ingest")

MAX_ROWS = 50_000
MAX_BYTES = 50 * 1024 * 1024
SNIFF_BYTES = 20_000
DELIMITERS = ",;\t|"
CONTENT_THRESHOLD = 0.6
SAMPLE_SIZE = 200

# Spec 14, in priority order.
HEADER_PRIORITY = (
    "website", "websiteurl", "companywebsite", "companywebsiteurl", "companydomain",
    "domain", "url", "companyurl", "homepage", "web", "site", "organizationwebsite",
    "businesswebsite",
)
EMAIL_HEADERS = ("email", "emailaddress", "contactemail", "workemail", "primaryemail",
                 "owneremail", "publicemail")

_EMAILISH = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,24}$", re.I)


class IngestError(ValueError):
    """`code` is the API error code (app/errors.py) this failure is reported as."""

    def __init__(self, message: str, code: str = "invalid_file") -> None:
        super().__init__(message)
        self.code = code


ALLOWED_EXTENSIONS = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm")


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[dict[str, str]]
    had_header: bool
    delimiter: str
    encoding: str

    def column(self, name: str) -> list[str]:
        return [(r.get(name) or "").strip() for r in self.rows]


@dataclass
class ColumnGuess:
    column: str | None = None
    confidence: str = "none"          # high | medium | low | none
    method: str = "none"              # header | content | email | none
    scores: dict[str, float] = field(default_factory=dict)
    email_column: str | None = None   # offered when no website column is found


def norm_header(name: str) -> str:
    return "".join(ch for ch in (name or "").lower() if ch.isalnum())


def decode(raw: bytes) -> tuple[str, str]:
    """BOM, then utf-8, then utf-16 when a BOM says so, then cp1252 (spec 14)."""
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        for enc in ("utf-16", "utf-16-le", "utf-16-be"):
            try:
                return raw.decode(enc), enc
            except (UnicodeDecodeError, LookupError):
                continue
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8/replace"


def sniff_delimiter(text: str, filename: str = "") -> str:
    if filename.lower().endswith(".tsv"):
        return "\t"
    head = text[:SNIFF_BYTES]
    try:
        return csv.Sniffer().sniff(head, delimiters=DELIMITERS).delimiter
    except csv.Error:
        # Fall back to whichever candidate appears most on the first line.
        first = head.splitlines()[0] if head.splitlines() else ""
        best = max(DELIMITERS, key=first.count)
        return best if first.count(best) else ","


def looks_like_header(cells: list[str]) -> bool:
    """Row 1 is a header unless every cell already parses as a domain or URL."""
    useful = [c for c in cells if c.strip()]
    if not useful:
        return True
    return not all(normalize_input(c).ok for c in useful)


def parse_bytes(raw: bytes, filename: str = "") -> ParsedFile:
    if len(raw) > MAX_BYTES:
        mb = 1024 * 1024
        raise IngestError(
            f"file is {len(raw) / mb:.1f} MB, limit is {MAX_BYTES // mb} MB",
            "file_too_large",
        )
    name = filename.lower()
    if name.endswith(".xls"):
        raise IngestError("old .xls files are not supported; save as .xlsx or .csv",
                          "unsupported_file_type")
    if "." in name.rsplit("/", 1)[-1] and not name.endswith(ALLOWED_EXTENSIONS):
        raise IngestError(
            f"{filename} is not a CSV, TSV, TXT or XLSX file", "unsupported_file_type"
        )

    if filename.lower().endswith((".xlsx", ".xlsm")):
        return _parse_xlsx(raw)

    text, encoding = decode(raw)
    delimiter = sniff_delimiter(text, filename)
    # The csv module handles quoted fields with embedded newlines and delimiters.
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    records = list(reader)
    if not records:
        raise IngestError("file is empty", "empty_file")

    first = [c.strip() for c in records[0]]
    had_header = looks_like_header(first)
    if had_header:
        headers = [h.strip() or f"column_{i + 1}" for i, h in enumerate(first)]
        body = records[1:]
    else:
        headers = [f"column_{i + 1}" for i in range(len(first))]
        body = records

    rows = [_to_row(headers, rec) for rec in body[:MAX_ROWS]]
    if len(body) > MAX_ROWS:
        raise IngestError(f"{len(body):,} data rows, limit is {MAX_ROWS:,}", "too_many_rows")
    return ParsedFile(headers, rows, had_header, delimiter, encoding)


def _to_row(headers: list[str], record: list[str]) -> dict[str, str]:
    row: dict[str, str] = {}
    for i, header in enumerate(headers):
        row[header] = record[i].strip() if i < len(record) else ""
    return row


def _parse_xlsx(raw: bytes) -> ParsedFile:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise IngestError("xlsx support requires openpyxl") from e

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - corrupt zip, not a workbook, encrypted
        raise IngestError("the spreadsheet could not be opened; re-save it as .xlsx or .csv",
                          "invalid_file") from e
    ws = wb[wb.sheetnames[0]]
    # Stop one past the cap: a small .xlsx can expand to millions of rows.
    records = []
    for row in ws.iter_rows(values_only=True):
        records.append(["" if c is None else str(c).strip() for c in row])
        if len(records) > MAX_ROWS + 1:
            break
    wb.close()
    if not records:
        raise IngestError("spreadsheet is empty", "empty_file")
    first = records[0]
    had_header = looks_like_header(first)
    headers = (
        [h or f"column_{i + 1}" for i, h in enumerate(first)] if had_header
        else [f"column_{i + 1}" for i in range(len(first))]
    )
    body = records[1:] if had_header else records
    if len(body) > MAX_ROWS:
        raise IngestError(f"{len(body):,} data rows, limit is {MAX_ROWS:,}", "too_many_rows")
    return ParsedFile(headers, [_to_row(headers, r) for r in body], had_header, ",", "xlsx")


def score_column(values: list[str]) -> float:
    """Share of non-empty values that normalise to a real, non-platform domain."""
    sample = [v for v in values if v.strip()][:SAMPLE_SIZE]
    if not sample:
        return 0.0
    good = 0
    for v in sample:
        n = normalize_input(v)
        if n.ok and n.domain_source == "website":
            good += 1
    return good / len(sample)


def score_emails(values: list[str]) -> float:
    sample = [v for v in values if v.strip()][:SAMPLE_SIZE]
    if not sample:
        return 0.0
    return sum(1 for v in sample if _EMAILISH.match(v)) / len(sample)


def social_share(values: list[str]) -> float:
    sample = [v for v in values if v.strip()][:SAMPLE_SIZE]
    if not sample:
        return 0.0
    bad = sum(1 for v in sample if normalize_input(v).reason == "social_or_directory")
    return bad / len(sample)


def detect_column(parsed: ParsedFile) -> ColumnGuess:
    normed = {norm_header(h): h for h in parsed.headers}

    # 1. header match
    for want in HEADER_PRIORITY:
        if want in normed:
            col = normed[want]
            return ColumnGuess(column=col, confidence="high", method="header")

    # 2. content scoring
    scores = {h: score_column(parsed.column(h)) for h in parsed.headers}
    ranked = sorted(
        parsed.headers,
        key=lambda h: (-scores[h], social_share(parsed.column(h)), parsed.headers.index(h)),
    )
    best = ranked[0] if ranked else None
    if best and scores[best] >= CONTENT_THRESHOLD:
        return ColumnGuess(column=best, confidence="medium", method="content", scores=scores)

    # 3. email fallback
    email_scores = {h: score_emails(parsed.column(h)) for h in parsed.headers}
    for want in EMAIL_HEADERS:
        if want in normed and email_scores.get(normed[want], 0) >= CONTENT_THRESHOLD:
            return ColumnGuess(
                column=None, confidence="low", method="email",
                scores=scores, email_column=normed[want],
            )
    best_email = max(email_scores, key=lambda h: email_scores[h], default=None)
    if best_email and email_scores[best_email] >= CONTENT_THRESHOLD:
        return ColumnGuess(
            column=None, confidence="low", method="email",
            scores=scores, email_column=best_email,
        )

    # 4. nothing found
    return ColumnGuess(scores=scores)


def normalize_column(
    parsed: ParsedFile, column: str, *, derive_from_email: bool = False
) -> list[NormalizedInput]:
    values = parsed.column(column)
    out: list[NormalizedInput] = []
    for v in values:
        n = normalize_input(v)
        # normalize_input already understands an email cell, including rejecting free
        # providers. Re-parsing the bare domain here would resurrect exactly those:
        # "bob@gmail.com" -> free_email_only -> "gmail.com" as a crawl target.
        if (
            derive_from_email
            and not n.ok
            and "@" in v
            and n.reason not in ("free_email_only", "social_or_directory")
        ):
            derived = normalize_input(v.split("@")[-1])
            if derived.ok:
                derived.domain_source = "email"
                n = derived
        out.append(n)
    return out


def preview(parsed: ParsedFile, guess: ColumnGuess, samples: int = 5,
            column: str | None = None) -> dict:
    """The payload POST /jobs/preview returns (spec 5).

    `column` previews a column the person picked instead of the detected one, so the
    sample table and counts follow the choice before a job starts. Domains are derived
    from emails under the same rule create_job uses: only for the offered email column.
    """
    column = column or guess.column or guess.email_column
    derive = guess.method == "email" and column == guess.email_column
    normalized = normalize_column(parsed, column, derive_from_email=derive) if column else []

    domains = [n.domain for n in normalized if n.ok]
    empties = sum(1 for n in normalized if n.reason == "empty")
    sample_rows = []
    for raw, n in list(zip(parsed.column(column) if column else [], normalized))[:samples]:
        sample_rows.append({
            "value": raw,
            "domain": n.domain,
            "status": "ok" if n.ok else "invalid_input",
            "reason": n.reason,
        })

    return {
        "column": column,
        "derived_from_email": derive,
        "detected_column": guess.column,
        "confidence": guess.confidence,
        "method": guess.method,
        "email_column": guess.email_column,
        "candidate_columns": [
            {"name": h, "score": round(guess.scores.get(h, 0.0), 3)} for h in parsed.headers
        ],
        "columns": parsed.headers,
        "row_count": len(parsed.rows),
        "empty_website_count": empties,
        "unique_domains": len(set(domains)),
        "duplicate_domains": len(domains) - len(set(domains)),
        "invalid_count": sum(1 for n in normalized if not n.ok),
        "samples": sample_rows,
        "free_email_skipped": sum(1 for n in normalized if n.reason == "free_email_only"),
        "had_header": parsed.had_header,
        "delimiter": parsed.delimiter,
        "encoding": parsed.encoding,
    }


__all__ = [
    "ColumnGuess", "IngestError", "ParsedFile", "detect_column", "normalize_column",
    "parse_bytes", "preview", "score_column", "score_emails", "FREE_EMAIL_LABELS",
]
